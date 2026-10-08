"""Bulk student upload: reading, checking and importing (roadmap 2.6, SPEC 5)."""

from __future__ import annotations

import datetime
import io

import pytest
from apps.academics.models import Batch
from apps.accounts.models import User
from apps.core.roles import Role
from apps.core.tenancy import TenantContextError, tenant_context
from apps.people.bulk_upload import (
    MAX_ROWS,
    TEMPLATE_COLUMNS,
    UploadError,
    build_template,
    check_rows,
    import_rows,
    read_rows,
)
from apps.people.models import GuardianStudentLink, StudentProfile
from openpyxl import load_workbook

from tests.conftest import TEST_INVALID_PASSWORD_FOR_VALIDATION, make_user
from tests.people.xlsx import row, xlsx


def _read(rows, **kw) -> list[dict]:
    return read_rows(io.BytesIO(xlsx(rows, **kw)))


def _check(school, *rows) -> list:
    institute = school.morning.institute
    with tenant_context(institute):
        return check_rows(institute, _read(list(rows)))


def _errors(school, **values) -> list[str]:
    return _check(school, row(**values))[0].errors


# Template and reading


def test_template_has_the_columns_and_a_help_sheet() -> None:
    workbook = load_workbook(io.BytesIO(build_template()))

    header = [cell.value for cell in workbook["Students"][1]]
    assert tuple(header) == TEMPLATE_COLUMNS
    assert workbook["Students"].max_row == 1
    assert "Help" in workbook.sheetnames


def test_read_skips_empty_rows_and_keeps_excel_row_numbers() -> None:
    rows = _read([row(), {}, row(student_email="b@example.com")])

    assert [r["number"] for r in rows] == [2, 4]
    assert rows[1]["values"]["student_email"] == "b@example.com"


def test_read_turns_numbers_and_dates_into_text() -> None:
    (only,) = _read([row(cnic=3520200000003, dob=datetime.datetime(2012, 5, 1))])

    assert only["values"]["cnic"] == "3520200000003"
    assert only["values"]["dob"] == "2012-05-01"


def test_read_accepts_header_in_any_case() -> None:
    header = [column.upper() for column in TEMPLATE_COLUMNS]
    workbook_rows = [{c.upper(): v for c, v in row().items()}]

    (only,) = _read(workbook_rows, header=header)

    assert only["values"]["full_name"] == "Alex Sample"


@pytest.mark.parametrize(
    ("data", "message"),
    [
        (b"not a zip file", "not a valid .xlsx workbook"),
        (xlsx([row()], header=("full_name", "batches")), "Missing columns:"),
        (xlsx([]), "no student rows"),
        (b"0" * (1024 * 1024 + 1), "larger than 1 MB"),
    ],
)
def test_read_rejects_bad_files(data, message) -> None:
    with pytest.raises(UploadError) as error:
        read_rows(io.BytesIO(data))

    assert message in str(error.value)


def test_read_limits_the_number_of_rows() -> None:
    many = [row(student_email=f"s{n}@example.com") for n in range(MAX_ROWS + 1)]

    with pytest.raises(UploadError) as error:
        _read(many)

    assert f"at most {MAX_ROWS}" in str(error.value)


# Checking rows


@pytest.mark.django_db
def test_good_row_is_ready(school) -> None:
    (check,) = _check(
        school,
        row(
            subjects="Morning:Math|Physics",
            class_label="grade 9",
            gender="Female",
            dob="2012-05-01",
        ),
    )

    assert check.ok, check.errors
    assert check.selections == {school.morning: [school.math, school.physics]}
    assert check.details.class_label == school.grade_9
    assert check.details.gender == "female"


@pytest.mark.django_db
@pytest.mark.parametrize(
    ("values", "message"),
    [
        ({"full_name": ""}, "Missing: full_name."),
        ({"batches": "Night", "subjects": "Night:Math"}, "Unknown or inactive batch"),
        ({"subjects": "Morning:Art"}, "Unknown or inactive subject: Art."),
        ({"batches": "Morning, Evening"}, "must match the batches column"),
        ({"subjects": "Morning:"}, "Choose at least one subject for Morning."),
        ({"dob": "01/05/2012"}, "Date of birth must be YYYY-MM-DD."),
        ({"gender": "x"}, "Gender must be Male, Female or Other."),
        ({"class_label": "Grade 1"}, "Unknown or inactive class: Grade 1."),
        ({"cnic": "123"}, "CNIC must be exactly 13 digits."),
        ({"student_email": "nope"}, "Student email is not a valid email address."),
        ({"guardian_email": "alex@example.com"}, "must be different"),
        ({"student_email": "t1-a@example.com"}, "already used by another account"),
        (
            {"student_password": TEST_INVALID_PASSWORD_FOR_VALIDATION},
            "Student password:",
        ),
        ({"guardian_password": ""}, "Enter a password for the new guardian."),
        ({"guardian_email": "t1-a@example.com"}, "Guardian email is already used"),
    ],
)
def test_row_errors(school, values, message) -> None:
    errors = _errors(school, **values)

    assert any(message in error for error in errors), errors


@pytest.mark.django_db
def test_inactive_batch_is_not_found(school) -> None:
    Batch.unscoped.filter(pk=school.morning.pk).update(is_active=False)

    assert "Unknown or inactive batch: Morning." in _errors(school)


@pytest.mark.django_db
def test_batch_of_other_institute_is_not_found(school) -> None:
    Batch.unscoped.filter(pk=school.batch_b.pk).update(name="Night")

    errors = _errors(school, batches="Night", subjects="Night:Math")

    assert "Unknown or inactive batch: Night." in errors


@pytest.mark.django_db
def test_existing_guardian_of_institute_needs_no_password(school) -> None:
    assert (
        _errors(school, guardian_email="G1-A@example.com", guardian_password="") == []
    )


@pytest.mark.django_db
def test_guardian_of_other_institute_is_rejected(school, institute_b) -> None:
    make_user(email="parent-b@example.com", role=Role.GUARDIAN, institute=institute_b)

    errors = _errors(school, guardian_email="parent-b@example.com")

    assert "Guardian email is already used by another account." in errors


@pytest.mark.django_db
def test_siblings_share_a_new_guardian(school) -> None:
    first, second = _check(
        school,
        row(),
        row(student_email="sibling@example.com", guardian_password=""),
    )

    assert first.ok and second.ok, (first.errors, second.errors)


@pytest.mark.django_db
def test_duplicates_inside_the_file(school) -> None:
    _, twice, crossed = _check(
        school,
        row(),
        row(guardian_email="other@example.com"),
        row(student_email="new@example.com", guardian_email="alex@example.com"),
    )

    assert "Student email is used twice in this file." in twice.errors
    assert "Guardian email is a student email in this file." in crossed.errors


@pytest.mark.django_db
def test_check_needs_tenant_context(school) -> None:
    with pytest.raises(TenantContextError):
        check_rows(school.morning.institute, _read([row()]))


# Import


@pytest.mark.django_db
def test_import_enrols_good_rows_and_skips_bad_ones(school) -> None:
    institute = school.morning.institute
    rows = _read(
        [
            row(),
            row(student_email="bad", guardian_email="x@example.com"),
            row(student_email="sibling@example.com", guardian_password=""),
        ]
    )
    with tenant_context(institute):
        result = import_rows(institute, rows)

    assert result.codes == ["STU-004", "STU-005"]
    assert [row.number for row in result.failed] == [3]
    # unscoped: test assertions outside a tenant context.
    guardians = {
        link.guardian_id
        for link in GuardianStudentLink.unscoped.filter(
            student__student_code__in=result.codes, institute=institute
        )
    }
    assert len(guardians) == 1


@pytest.mark.django_db
def test_import_checks_rows_again(school) -> None:
    institute = school.morning.institute
    rows = _read([row()])
    make_user(email="alex@example.com", role=Role.STUDENT, institute=institute)

    with tenant_context(institute):
        result = import_rows(institute, rows)

    assert result.codes == []
    assert "Student email is already used by another account." in (
        result.failed[0].errors
    )
    # unscoped: test assertion outside a tenant context.
    assert not StudentProfile.unscoped.filter(user__email="alex@example.com").exists()


@pytest.mark.django_db
def test_import_needs_tenant_context(school) -> None:
    with pytest.raises(TenantContextError):
        import_rows(school.morning.institute, _read([row()]))
    assert not User.objects.filter(email="alex@example.com").exists()


@pytest.mark.django_db
def test_import_skips_rows_that_only_the_check_rejects(school) -> None:
    # enrol_student would accept a missing date or crash on a missing name,
    # so import_rows must skip every row the check marks as bad.
    institute = school.morning.institute
    rows = _read([row(dob="01/05/2012"), row(full_name="", student_email="b@x.com")])

    with tenant_context(institute):
        result = import_rows(institute, rows)

    assert result.codes == []
    assert [row.number for row in result.failed] == [2, 3]
    assert not User.objects.filter(email__in=["alex@example.com", "b@x.com"]).exists()
